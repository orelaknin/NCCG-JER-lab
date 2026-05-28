from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from openpyxl import load_workbook

try:
    from .ollama_utils import ExtractionResult, extract_temperature_info
    from .pdf_utils import extract_pdf_text
    from .path_utils import connect_to_share, normalize_datasheet_path
except ImportError:
    from ollama_utils import ExtractionResult, extract_temperature_info
    from pdf_utils import extract_pdf_text
    from path_utils import connect_to_share, normalize_datasheet_path


@dataclass
class ProcessingConfig:
    input_excel: str
    output_excel: str
    datasheet_column: str
    unique_id_column: str
    drive_letter: str = ""
    network_root: str = ""
    username: str = ""
    password: str = ""
    domain: str = ""
    model: str = "llama3.1"
    prompt: str = ""


ProgressCallback = Callable[[int, int], None]
LogCallback = Callable[[str], None]


DEFAULT_OUTPUT_COLUMNS = {
    "normalized_datasheet_path": "normalized_datasheet_path",
    "temperature_range": "temperature_range",
    "source_sentence": "source_sentence",
    "extraction_status": "extraction_status",
    "extraction_error": "extraction_error",
}


def load_excel_columns(excel_path: str) -> list[str]:
    """Return column names from the first sheet with a fast header-only read."""

    suffix = Path(excel_path).suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if header_row is None:
                return []
            return [str(cell).strip() for cell in header_row if cell is not None and str(cell).strip()]
        finally:
            workbook.close()

    preview = pd.read_excel(excel_path, nrows=0, engine="openpyxl")
    return [str(col).strip() for col in preview.columns.tolist() if str(col).strip()]


def _safe_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def process_excel(
    config: ProcessingConfig,
    progress_callback: Optional[ProgressCallback] = None,
    log_callback: Optional[LogCallback] = None,
) -> str:
    """Process the workbook and create a new Excel file with extracted data."""

    def log(message: str) -> None:
        if log_callback:
            log_callback(message)

    def progress(current: int, total: int) -> None:
        if progress_callback:
            progress_callback(current, total)

    input_path = Path(config.input_excel)
    output_path = Path(config.output_excel)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    log(f"Reading workbook: {input_path}")
    df = pd.read_excel(input_path, engine="openpyxl")
    if config.datasheet_column not in df.columns:
        raise ValueError(f"Datasheet column '{config.datasheet_column}' was not found in the workbook.")
    if config.unique_id_column not in df.columns:
        raise ValueError(f"Unique ID column '{config.unique_id_column}' was not found in the workbook.")

    if config.network_root:
        ok, message = connect_to_share(config.network_root, config.username, config.password, config.domain)
        log(message)
        if not ok:
            log("Continuing without authenticated share connection.")

    total = len(df)
    log(f"Rows to process: {total}")

    normalized_paths: list[str] = []
    temp_ranges: list[str] = []
    source_sentences: list[str] = []
    statuses: list[str] = []
    errors: list[str] = []

    task_prompt = config.prompt.strip()
    if not task_prompt:
        task_prompt = (
            "Extract the operating temperature range from the datasheet text and return the exact source sentence."
        )

    for idx, row in df.iterrows():
        uid = _safe_text(row.get(config.unique_id_column, ""))
        raw_path = row.get(config.datasheet_column, "")
        normalized_path = normalize_datasheet_path(raw_path, config.drive_letter, config.network_root)
        normalized_paths.append(normalized_path)

        current_row = len(normalized_paths)
        progress(current_row, total)

        if not normalized_path:
            temp_ranges.append("")
            source_sentences.append("")
            statuses.append("missing_path")
            errors.append("Datasheet path is empty")
            log(f"[{current_row}/{total}] {uid or idx}: missing datasheet path")
            continue

        if not Path(normalized_path).exists():
            temp_ranges.append("")
            source_sentences.append("")
            statuses.append("file_not_found")
            errors.append(f"File not found: {normalized_path}")
            log(f"[{current_row}/{total}] {uid or idx}: file not found -> {normalized_path}")
            continue

        pdf_result = extract_pdf_text(normalized_path)
        if pdf_result.error:
            temp_ranges.append("")
            source_sentences.append("")
            statuses.append("pdf_read_error")
            errors.append(pdf_result.error)
            log(f"[{current_row}/{total}] {uid or idx}: {pdf_result.error}")
            continue

        if not pdf_result.text.strip():
            temp_ranges.append("")
            source_sentences.append("")
            statuses.append("no_text_extracted")
            errors.append("No searchable text extracted from PDF")
            log(f"[{current_row}/{total}] {uid or idx}: no searchable text extracted")
            continue

        extraction: ExtractionResult = extract_temperature_info(
            pdf_result.text,
            prompt=task_prompt,
            model=config.model,
        )

        temp_ranges.append(extraction.temperature_range)
        source_sentences.append(extraction.source_sentence)
        if extraction.error:
            statuses.append("llm_warning")
            errors.append(extraction.error)
            log(f"[{current_row}/{total}] {uid or idx}: LLM warning -> {extraction.error}")
        else:
            statuses.append("ok")
            errors.append("")
            log(f"[{current_row}/{total}] {uid or idx}: extracted temperature range")

    df[DEFAULT_OUTPUT_COLUMNS["normalized_datasheet_path"]] = normalized_paths
    df[DEFAULT_OUTPUT_COLUMNS["temperature_range"]] = temp_ranges
    df[DEFAULT_OUTPUT_COLUMNS["source_sentence"]] = source_sentences
    df[DEFAULT_OUTPUT_COLUMNS["extraction_status"]] = statuses
    df[DEFAULT_OUTPUT_COLUMNS["extraction_error"]] = errors

    log(f"Writing output workbook: {output_path}")
    df.to_excel(output_path, index=False, engine="openpyxl")
    return str(output_path)
